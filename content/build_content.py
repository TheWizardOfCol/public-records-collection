#!/usr/bin/env python3
"""Build the Public Records Collection content data from the raw source files.

Reads the curated spreadsheets/CSV (books, prompts, parody titles, sticker
prompts) and writes normalized JSON under ``content/generated/``. Book text and
sticker images are handled separately (see ``content/README.md``) so large raw
assets never enter this repository.

Usage:
    python3 content/build_content.py [--source /workspace/Prompts] [--out content/generated]

Requires openpyxl (``pip install --user --break-system-packages openpyxl``).
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import unicodedata

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install --user --break-system-packages openpyxl")


def truthy(value) -> bool:
    """Normalize the spreadsheet's t/f truth flags."""
    return str(value).strip().lower() in {"t", "true", "1", "yes", "y"}


def slugify(text: str) -> str:
    """Lowercase, strip accents/possessives/punctuation — matches book files to titles."""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.lower().replace("&", " and ")
    text = re.sub(r"[’'`]s\b", "s", text)          # possessives
    text = re.sub(r"[’'`]", "", text)
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


# Book files whose name doesn't survive slugification cleanly (encoding or
# punctuation quirks in the source drop). Keyed by normalized title slug.
BOOK_FILE_ALIASES = {
    "lesmiserables": "Les Mis+¬rables.txt",
    "theartandpracticeoftypography": "The Art  Practice of Typography.txt",
}


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def world_id(prefix: str, index: int) -> str:
    return f"{prefix}-{index:04d}"


def read_rows(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb


def build_books(source: str, books_dir: str):
    wb = read_rows(os.path.join(source, "get-lit-content", "Get Lit Content.xlsx"))
    ws = wb["Books"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    files_by_slug = {}
    if os.path.isdir(books_dir):
        for name in os.listdir(books_dir):
            files_by_slug[slugify(os.path.splitext(name)[0])] = name

    books = []
    for i, r in enumerate(rows, start=1):
        title = clean(r[0])
        if not title:
            continue
        slug = slugify(title)
        books.append({
            "id": world_id("book", i),
            "title": title,
            "author": clean(r[1]),
            "category": clean(r[2]),
            "description": clean(r[3]),
            "triggerWarnings": clean(r[4]),
            "isSpicy": truthy(r[5]),
            "isActive": truthy(r[6]),
            "slug": slug,
            "textFile": files_by_slug.get(slug) or BOOK_FILE_ALIASES.get(slug),
        })
    return books


def build_prompts(source: str):
    wb = read_rows(os.path.join(source, "get-lit-content", "Get Lit Content.xlsx"))
    ws = wb["Prompts"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    prompts = []
    for i, r in enumerate(rows, start=1):
        text = clean(r[0])
        if not text:
            continue
        prompts.append({
            "id": world_id("prompt", i),
            "text": text,
            "collection": clean(r[1]),
            "deck": clean(r[2]),
            "isSpicy": truthy(r[3]),
            "isActive": truthy(r[4]),
        })
    return prompts


def build_parody_titles(source: str):
    wb = read_rows(os.path.join(source, "get-lit-content", "Get Lit Content.xlsx"))
    ws = wb["Parody Titles"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    titles = []
    for i, r in enumerate(rows, start=1):
        title = clean(r[0])
        if not title:
            continue
        titles.append({"id": world_id("parody", i), "title": title})
    return titles


def build_sticker_prompts(source: str):
    path = os.path.join(source, "sticker-doodle-prompts.csv")
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    prompts = []
    for i, r in enumerate(rows, start=1):
        text = clean(r.get("text"))
        if not text:
            continue
        prompts.append({
            "id": world_id("sticker", i),
            "text": text,
            "category": clean(r.get("category")),
            "isActive": truthy(r.get("is_active")),
        })
    return prompts


def write_json(path: str, payload):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"wrote {path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--source", default="/workspace/Prompts")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "generated"))
    args = ap.parse_args()

    books_dir = os.path.join(args.source, "get-lit-content", "books")
    out = args.out

    books = build_books(args.source, books_dir)
    prompts = build_prompts(args.source)
    parody = build_parody_titles(args.source)
    stickers = build_sticker_prompts(args.source)

    write_json(os.path.join(out, "get-lit", "books.json"), books)
    write_json(os.path.join(out, "get-lit", "prompts.json"), prompts)
    write_json(os.path.join(out, "get-lit", "parody_titles.json"), parody)
    write_json(os.path.join(out, "sticker-doodle", "prompts.json"), stickers)

    mapped = sum(1 for b in books if b["textFile"])
    print(f"\nbooks: {len(books)} ({mapped} with text file, {len(books) - mapped} unmapped)  "
          f"active={sum(b['isActive'] for b in books)} spicy={sum(b['isSpicy'] for b in books)}")
    print(f"prompts: {len(prompts)}  parody titles: {len(parody)}  sticker prompts: {len(stickers)}")


if __name__ == "__main__":
    main()
